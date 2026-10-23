from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Iterable

if TYPE_CHECKING:
    from langchain_core.documents import Document

SUPPORTED = {".pdf", ".txt", ".docx", ".md"}

def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader

        return PyPDFLoader(str(path)).load()

    if suffix in {".txt", ".md"}:
        from langchain_community.document_loaders import TextLoader

        return TextLoader(str(path), encoding="utf-8").load()

    if suffix == ".docx":
        from docx import Document as DocxDocument
        from langchain_core.documents import Document

        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]

    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs: Iterable[Document]) -> list[Document]:
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)
    return splitter.split_documents(list(docs))
